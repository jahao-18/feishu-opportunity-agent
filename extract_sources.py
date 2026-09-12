from pathlib import Path
from docx import Document
import pdfplumber


DOCX = Path(r"C:\Users\18304\Desktop\FDE 实习生考试题库(1).docx")
PDF = Path(r"C:\Users\18304\Desktop\08题 澄初个人护理.pdf")


def extract_docx(path: Path) -> str:
    doc = Document(path)
    chunks: list[str] = []
    for index, paragraph in enumerate(doc.paragraphs, 1):
        text = paragraph.text.strip()
        if text:
            chunks.append(f"P{index}: {text}")
    for table_index, table in enumerate(doc.tables, 1):
        chunks.append(f"\n[TABLE {table_index}]")
        for row_index, row in enumerate(table.rows, 1):
            cells = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
            chunks.append(f"R{row_index}: " + " | ".join(cells))
    return "\n".join(chunks)


def extract_pdf(path: Path) -> str:
    chunks: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page_index, page in enumerate(pdf.pages, 1):
            chunks.append(f"\n===== PDF PAGE {page_index} =====\n")
            chunks.append(page.extract_text(x_tolerance=2, y_tolerance=3) or "[NO TEXT]")
    return "\n".join(chunks)


if __name__ == "__main__":
    print("===== DOCX =====")
    print(extract_docx(DOCX))
    print("\n===== PDF =====")
    print(extract_pdf(PDF))
